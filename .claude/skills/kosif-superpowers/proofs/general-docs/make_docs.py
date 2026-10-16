#!/usr/bin/env python3
"""general-docs proof: write Arabic+English DOCX, XLSX, PPTX and PDF, convert them
with LibreOffice and pandoc, then read every file back and check the Arabic survived.

Usage:  python -I make_docs.py <workdir>
Needs:  python-docx python-pptx openpyxl reportlab weasyprint pypdf pdfplumber pypdfium2
        arabic-reshaper python-bidi "markitdown[docx,pdf,xlsx,pptx]" pillow pytesseract
        + soffice, pdftotext, pdffonts, pandoc, tesseract (+ara) on PATH. Fonts come from ../_shared/fonts (OFL).
Writes: <workdir>/... (all generated files), results.json and contact.png next to this script.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import time
import unicodedata
from pathlib import Path

HERE = Path(__file__).resolve().parent
FONTS = HERE.parent / "_shared" / "fonts"
W = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/general-docs").resolve()
W.mkdir(parents=True, exist_ok=True)

TITLE_AR = "تقرير المبيعات الربعي"
LINE_AR = "ارتفعت المبيعات بنسبة ١٢٫٥٪ في الربع الثالث."
LINE_EN = "Sales rose 12.5% in Q3 (mixed RTL + LTR line)."
ROWS = [("القاهرة", 1200), ("الرياض", 950), ("دبي", 780), ("عمّان", 410)]
KEY = "المبيعات"  # word that must survive every round-trip

R: dict = {"workdir": str(W), "checks": [], "timings_s": {}}


def check(name: str, ok: bool, detail="") -> None:
    R["checks"].append({"name": name, "ok": bool(ok), "detail": detail})
    print(("PASS " if ok else "FAIL ") + name + (f" | {detail}" if detail else ""))


def has_ar(text: str, word: str = KEY) -> bool:
    """True when the word appears in logical order after NFKC (undoes presentation forms)."""
    return word in unicodedata.normalize("NFKC", text or "")


def run(cmd: list[str], env=None, timeout=180) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, capture_output=True, text=True, env=env, timeout=timeout)


# A private fontconfig so LibreOffice/pango see the OFL Arabic fonts without touching ~/.fonts
FC = W / "fonts.conf"
FC.write_text(f"""<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>
<include ignore_missing="yes">/etc/fonts/fonts.conf</include><dir>{FONTS}</dir>
<cachedir>{W / 'fc-cache'}</cachedir></fontconfig>""")
ENV = dict(os.environ, FONTCONFIG_FILE=str(FC))
os.environ["FONTCONFIG_FILE"] = str(FC)  # weasyprint (pango) in this process too

# ---------------------------------------------------------------- DOCX (python-docx)
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def rtl_paragraph(p, font="Amiri", size=14):
    ppr = p._p.get_or_add_pPr()
    ppr.append(OxmlElement("w:bidi"))  # paragraph direction RTL
    for run_ in p.runs:
        rpr = run_._r.get_or_add_rPr()
        rpr.append(OxmlElement("w:rtl"))  # run is right-to-left
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:  # lxml elements without children are falsy: never use `or` here
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        fonts.set(qn("w:cs"), font)  # complex-script font is what Arabic uses
        fonts.set(qn("w:ascii"), font)
        fonts.set(qn("w:hAnsi"), font)
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(size * 2))
        rpr.append(szcs)
        run_.font.size = Pt(size)


t0 = time.time()
d = docx.Document()
rtl_paragraph(d.add_heading(TITLE_AR, level=1), size=20)
rtl_paragraph(d.add_paragraph(LINE_AR))
d.add_paragraph(LINE_EN)
tbl = d.add_table(rows=1, cols=2)
tbl.style = "Table Grid"
tblpr = tbl._tbl.tblPr
tblpr.append(OxmlElement("w:bidiVisual"))  # first column on the right
for cell, txt in zip(tbl.rows[0].cells, ["المدينة", "المبيعات"]):
    cell.text = txt
    rtl_paragraph(cell.paragraphs[0], size=12)
for city, val in ROWS:
    c = tbl.add_row().cells
    c[0].text, c[1].text = city, f"{val:,}"
    rtl_paragraph(c[0].paragraphs[0], size=12)
DOCX = W / "report_ar.docx"
d.save(DOCX)
back = docx.Document(DOCX)
txt = "\n".join(p.text for p in back.paragraphs) + "\n" + "\n".join(c.text for r in back.tables[0].rows for c in r.cells)
import zipfile
docxml = zipfile.ZipFile(DOCX).read("word/document.xml").decode()
check("docx: write + read back Arabic heading/table", TITLE_AR in txt and "عمّان" in txt, f"{DOCX.stat().st_size} B")
check("docx: RTL markers present (w:bidi, w:rtl, w:bidiVisual)", all(k in docxml for k in ("<w:bidi/>", "<w:rtl/>", "<w:bidiVisual/>")))
R["timings_s"]["docx"] = round(time.time() - t0, 2)

# ---------------------------------------------------------------- XLSX (openpyxl)
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font

t0 = time.time()
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "المبيعات"
ws.sheet_view.rightToLeft = True  # sheet opens right-to-left (column A on the right)
ws.append(["المدينة", "المبيعات", "الحصة"])
for city, val in ROWS:
    ws.append([city, val, None])
n = len(ROWS) + 1
for r in range(2, n + 1):
    ws.cell(r, 3, f"=B{r}/SUM($B$2:$B${n})").number_format = "0.0%"
ws.cell(n + 1, 1, "الإجمالي")
ws.cell(n + 1, 2, f"=SUM(B2:B{n})").number_format = "#,##0"
for c in ws[1]:
    c.font = Font(bold=True, name="Cairo")
    c.alignment = Alignment(horizontal="right", readingOrder=2)  # 2 = RTL reading order
ch = BarChart()
ch.title = "Sales by city"
ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=n), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
ws.add_chart(ch, "E2")
ws.page_setup.orientation = "landscape"  # without fit-to-page LibreOffice splits the chart over 2 PDF pages
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
XLSX = W / "sales_ar.xlsx"
wb.save(XLSX)
raw = openpyxl.load_workbook(XLSX)["المبيعات"]
check("xlsx: openpyxl stores formulas as text, no cached value",
      str(raw.cell(n + 1, 2).value).startswith("=SUM") and openpyxl.load_workbook(XLSX, data_only=True)["المبيعات"].cell(n + 1, 2).value is None,
      f"B{n+1}={raw.cell(n + 1, 2).value!r}; data_only -> None")
check("xlsx: sheet is RTL and has a native bar chart", raw.sheet_view.rightToLeft is True and len(raw._charts) == 1)
# Recalculate with LibreOffice: xlsx -> xlsx round trip fills the cached values
RECALC = W / "recalc"
RECALC.mkdir(exist_ok=True)
t1 = time.time()
p = run(["soffice", "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", str(RECALC), str(XLSX)], env=ENV)
R["timings_s"]["soffice_xlsx_recalc"] = round(time.time() - t1, 2)
vals = openpyxl.load_workbook(RECALC / XLSX.name, data_only=True).worksheets[0]
total = vals.cell(n + 1, 2).value
share = vals.cell(2, 3).value
check("xlsx: LibreOffice headless recalculates formulas",
      total == sum(v for _, v in ROWS) and abs((share or 0) - 1200 / 3340) < 1e-9,
      f"total={total}, share(القاهرة)={share}")
R["timings_s"]["xlsx"] = round(time.time() - t0, 2)

# ---------------------------------------------------------------- PPTX (python-pptx)
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt as PPt
from pptx.oxml.ns import qn as pqn

t0 = time.time()
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9


def rtl_text_frame(tf, size=32, font="Cairo"):
    for para in tf.paragraphs:
        ppr = para._p.get_or_add_pPr()
        ppr.set("rtl", "1")
        ppr.set("algn", "r")
        for run_ in para.runs:
            run_.font.size = PPt(size)
            rpr = run_._r.get_or_add_rPr()
            rpr.set("lang", "ar-SA")
            for tag in ("a:latin", "a:cs"):  # a:cs = complex-script typeface (Arabic)
                el = rpr.find(pqn(tag))
                if el is None:
                    el = rpr.makeelement(pqn(tag), {})
                    rpr.append(el)
                el.set("typeface", font)


s1 = prs.slides.add_slide(prs.slide_layouts[1])
s1.shapes.title.text = TITLE_AR
rtl_text_frame(s1.shapes.title.text_frame, 40)
body = s1.placeholders[1].text_frame
body.text = LINE_AR
body.add_paragraph().text = "القاهرة في المركز الأول"
rtl_text_frame(body, 24)
s2 = prs.slides.add_slide(prs.slide_layouts[5])
s2.shapes.title.text = "المبيعات حسب المدينة"
rtl_text_frame(s2.shapes.title.text_frame, 36)
cd = CategoryChartData()
cd.categories = [c for c, _ in ROWS]
cd.add_series("Sales", [v for _, v in ROWS])
s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(11.3), Inches(5.4), cd)
PPTX = W / "deck_ar.pptx"
prs.save(PPTX)
pb = Presentation(PPTX)
slide_txt = [sh.text_frame.text for sl in pb.slides for sh in sl.shapes if sh.has_text_frame]
chart = [sh for sh in pb.slides[1].shapes if sh.has_chart][0].chart
check("pptx: write + read back Arabic title/body", TITLE_AR in slide_txt and any(KEY in t for t in slide_txt), f"{len(pb.slides)} slides, {PPTX.stat().st_size} B")
check("pptx: native chart categories read back", list(chart.plots[0].categories) == [c for c, _ in ROWS], str(list(chart.plots[0].categories)))
R["timings_s"]["pptx"] = round(time.time() - t0, 2)

# ---------------------------------------------------------------- PDF (weasyprint, HTML/CSS)
import weasyprint

t0 = time.time()
rows_html = "".join(f"<tr><td>{c}</td><td>{v:,}</td></tr>" for c, v in ROWS)
html = f"""<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8"><style>
@font-face {{ font-family: Amiri; src: url('{(FONTS / 'Amiri-Regular.ttf').as_uri()}'); }}
@font-face {{ font-family: Cairo; src: url('{(FONTS / 'Cairo[slnt,wght].ttf').as_uri()}'); }}
@page {{ size: A4; margin: 16mm; }}
body {{ font-family: Amiri, serif; font-size: 15pt; }}
h1 {{ font-family: Cairo; font-weight: 800; color: #1d3b6e; margin: 0 0 6mm; }}
table {{ border-collapse: collapse; width: 70%; }} td {{ border: 1px solid #999; padding: 2mm 4mm; }}
.en {{ direction: ltr; text-align: left; font-family: Cairo; font-size: 11pt; }}
</style></head><body><h1>{TITLE_AR}</h1><p>{LINE_AR}</p><p class="en">{LINE_EN}</p>
<table><tr><td><b>المدينة</b></td><td><b>المبيعات</b></td></tr>{rows_html}</table></body></html>"""
PDF = W / "report_weasy.pdf"
weasyprint.HTML(string=html, base_url=str(W)).write_pdf(PDF)
R["timings_s"]["weasyprint"] = round(time.time() - t0, 2)

import pdfplumber
import pypdf

plumb = pdfplumber.open(PDF).pages[0].extract_text() or ""
pyp = pypdf.PdfReader(PDF).pages[0].extract_text() or ""
check("pdf(weasyprint): pdfplumber/pdfminer return Arabic REVERSED (visual order) -> use pypdf or pdftotext",
      not has_ar(plumb) and has_ar(plumb[::-1]), repr(plumb.splitlines()[0][:60]) if plumb else "empty")
check("pdf(weasyprint): Arabic recoverable with pypdf", has_ar(pyp), repr(pyp.splitlines()[0][:60]) if pyp else "empty")
pt = run(["pdftotext", "-layout", str(PDF), "-"]).stdout
check("pdf(weasyprint): Arabic recoverable with pdftotext", has_ar(pt), repr(pt.strip().splitlines()[0][:60]) if pt.strip() else "empty")
fonts_in_pdf = sorted({f.get("/BaseFont", "") for pg in pypdf.PdfReader(PDF).pages for f in (pg["/Resources"].get("/Font") or {}).values() for f in [f.get_object()]})
check("pdf(weasyprint): Arabic fonts embedded as subsets", any("Amiri" in f for f in fonts_in_pdf), ", ".join(fonts_in_pdf))

# ---------------------------------------------------------------- PDF (reportlab: needs shaping by hand)
import arabic_reshaper
from bidi.algorithm import get_display
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

pdfmetrics.registerFont(TTFont("Amiri", str(FONTS / "Amiri-Regular.ttf")))
RL = W / "report_reportlab.pdf"
c = canvas.Canvas(str(RL), pagesize=(595, 300))
c.setFont("Amiri", 22)
c.drawRightString(560, 250, TITLE_AR)  # raw: letters unjoined and reversed
c.drawRightString(560, 200, get_display(arabic_reshaper.reshape(TITLE_AR)))  # shaped + visual order
c.setFont("Amiri", 12)
c.drawRightString(560, 160, "line 1 = raw string (broken); line 2 = arabic_reshaper + python-bidi (correct)")
c.save()
rl_txt = pdfplumber.open(RL).pages[0].extract_text() or ""
check("pdf(reportlab): shaped line written (reshaper+bidi needed)", RL.stat().st_size > 0, f"{RL.stat().st_size} B; extract={rl_txt.splitlines()[1][:40]!r}" if len(rl_txt.splitlines()) > 1 else rl_txt[:40])

# ---------------------------------------------------------------- LibreOffice: docx/pptx/xlsx -> PDF
CONV = W / "soffice_pdf"
CONV.mkdir(exist_ok=True)
t0 = time.time()
p = run(["soffice", "--headless", "--norestore", "--convert-to", "pdf", "--outdir", str(CONV), str(DOCX), str(PPTX), str(XLSX)], env=ENV, timeout=300)
R["timings_s"]["soffice_3_to_pdf"] = round(time.time() - t0, 2)
for src in (DOCX, PPTX, XLSX):
    out = CONV / (src.stem + ".pdf")
    t = run(["pdftotext", str(out), "-"]).stdout if out.exists() else ""
    word = TITLE_AR if src is not XLSX else "القاهرة"
    check(f"soffice {src.suffix}->pdf: file made and Arabic text recoverable", out.exists() and has_ar(t, word), f"{out.stat().st_size if out.exists() else 0} B")
pf = run(["pdffonts", str(CONV / "report_ar.pdf")]).stdout
check("soffice docx->pdf: uses the requested Amiri font (private fontconfig works)", "Amiri" in pf, " ".join(l.split()[0] for l in pf.splitlines()[2:]))

# ---------------------------------------------------------------- markitdown: everything -> Markdown
from markitdown import MarkItDown

md = MarkItDown()
for f in (DOCX, XLSX, PPTX, PDF):
    t0 = time.time()
    out = md.convert(str(f)).text_content
    (W / (f.name + ".md")).write_text(out, encoding="utf-8")
    word = KEY if f is not XLSX else "القاهرة"
    if f is PDF:  # markitdown reads PDFs through pdfminer, so Arabic comes out reversed too
        check("markitdown .pdf -> md: Arabic comes out REVERSED (pdfminer) -> prefer pypdf/pdftotext for RTL PDFs",
              not has_ar(out, word) and has_ar(out[::-1], word), f"{len(out)} chars in {time.time() - t0:.2f}s")
        continue
    check(f"markitdown {f.suffix} -> md keeps Arabic", has_ar(out, word), f"{len(out)} chars in {time.time() - t0:.2f}s")

# ---------------------------------------------------------------- pandoc: Markdown <-> DOCX
MD = W / "note.md"
MD.write_text(f"---\ndir: rtl\nlang: ar\n---\n\n# {TITLE_AR}\n\n{LINE_AR}\n\n| المدينة | المبيعات |\n|---|---|\n" + "".join(f"| {c} | {v} |\n" for c, v in ROWS), encoding="utf-8")
PDX = W / "note_pandoc.docx"
p1 = run(["pandoc", str(MD), "-o", str(PDX)])
p2 = run(["pandoc", str(PDX), "-t", "gfm"])
check("pandoc md->docx->md round-trip keeps Arabic table", p1.returncode == 0 and "عمّان" in p2.stdout and TITLE_AR in p2.stdout, f"{PDX.stat().st_size} B")

# ---------------------------------------------------------------- pypdf: merge, rotate, encrypt
w = pypdf.PdfWriter()
for f in (PDF, CONV / "report_ar.pdf", CONV / "deck_ar.pdf"):
    w.append(str(f))
w.pages[0].rotate(90)
w.encrypt(user_password="kosif", algorithm="AES-256")
BUNDLE = W / "bundle_locked.pdf"
w.write(BUNDLE)
rd = pypdf.PdfReader(BUNDLE)
locked = rd.is_encrypted
rd.decrypt("kosif")
check("pypdf: merge 3 PDFs, rotate p1, AES-256 encrypt, reopen with password", locked and len(rd.pages) >= 4 and rd.pages[0].get("/Rotate") == 90, f"{len(rd.pages)} pages")

# ---------------------------------------------------------------- "scanned" PDF -> OCR (pypdfium2 render + tesseract)
import pypdfium2 as pdfium
import pytesseract

scan = pdfium.PdfDocument(str(PDF))[0].render(scale=300 / 72).to_pil().convert("L")  # 300 dpi raster = no text layer
scan.save(W / "scan_300dpi.png")
t0 = time.time()
ocr = pytesseract.image_to_string(scan, lang="ara+eng", config="--oem 1 --psm 4")
R["timings_s"]["ocr_page_300dpi"] = round(time.time() - t0, 2)
check("scanned page (300 dpi raster of the PDF) -> tesseract ara+eng recovers title + English line",
      has_ar(ocr, TITLE_AR) and "Sales rose 12.5%" in ocr, f"{len(ocr)} chars in {R['timings_s']['ocr_page_300dpi']} s")

# ---------------------------------------------------------------- contact sheet (pypdfium2 render)
from PIL import Image, ImageChops, ImageDraw, ImageFont

tiles = []
for label, f in [("weasyprint HTML->PDF", PDF), ("python-docx -> soffice PDF", CONV / "report_ar.pdf"),
                 ("python-pptx -> soffice PDF (slide 1)", CONV / "deck_ar.pdf"), ("python-pptx chart (slide 2)", CONV / "deck_ar.pdf"),
                 ("openpyxl RTL sheet -> soffice PDF", CONV / "sales_ar.pdf"), ("reportlab raw vs reshaped", RL)]:
    doc = pdfium.PdfDocument(str(f))
    page = doc[1 if "slide 2" in label else 0]
    im = page.render(scale=2).to_pil().convert("RGB")
    box = ImageChops.difference(im, Image.new("RGB", im.size, "white")).getbbox()  # crop to the ink
    if box:
        im = im.crop((max(box[0] - 20, 0), max(box[1] - 20, 0), min(box[2] + 20, im.width), min(box[3] + 20, im.height)))
    im.thumbnail((560, 400))
    tile = Image.new("RGB", (580, 440), "white")
    tile.paste(im, ((580 - im.width) // 2, 34))
    ImageDraw.Draw(tile).text((10, 8), label, fill="#b00020", font=ImageFont.truetype(str(FONTS / "Cairo[slnt,wght].ttf"), 18))
    tiles.append(tile)
sheet = Image.new("RGB", (580 * 3, 440 * 2), "#dddddd")
for i, t in enumerate(tiles):
    sheet.paste(t, ((i % 3) * 580, (i // 3) * 440))
sheet = sheet.quantize(colors=128, method=Image.Quantize.MEDIANCUT)
sheet.save(HERE / "contact.png", optimize=True)
check("contact.png written", (HERE / "contact.png").stat().st_size < 300_000, f"{(HERE / 'contact.png').stat().st_size} B")

# ---------------------------------------------------------------- versions + summary
from importlib.metadata import version as v
R["versions"] = {k: v(k) for k in ("python-docx", "python-pptx", "openpyxl", "reportlab", "weasyprint", "pypdf", "pdfplumber", "pypdfium2", "markitdown", "arabic-reshaper", "python-bidi", "pytesseract")}
R["versions"]["tesseract"] = str(pytesseract.get_tesseract_version())
R["versions"]["soffice"] = run(["soffice", "--version"]).stdout.strip()
R["versions"]["pandoc"] = run(["pandoc", "--version"]).stdout.splitlines()[0]
R["sizes_B"] = {f.name: f.stat().st_size for f in (DOCX, XLSX, PPTX, PDF, RL, PDX, BUNDLE)}
ok = sum(c["ok"] for c in R["checks"])
R["summary"] = f"{ok}/{len(R['checks'])} checks passed"
print(R["summary"])
(HERE / "results.json").write_text(json.dumps(R, ensure_ascii=False, indent=1), encoding="utf-8")
